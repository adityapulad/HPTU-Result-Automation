"""
HPTU Result Scraper
Scrapes B.TECH CSE (AI-DS) 4th Semester Results (May 2026) from HPTU portal
and compiles them into an Excel spreadsheet.

Default Roll Numbers: 240505010001 to 240505010075 (+ extras)

Usage:
    python scrape_results.py
    python scrape_results.py --url "https://...query.aspx?id=..." --start 240505010001 --end 240505010075
    python scrape_results.py --skip 240505010043 --extra 240501010267 240507010176
"""

import argparse
import re
import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import time
import sys

from api.result import _parse_structured

# Configuration (defaults; override via CLI flags)
BASE_URL = "https://himturesult.indiaresults.com/hp/himtu/hp-himtu"
QUERY_URL = f"{BASE_URL}/query.aspx?id=1800266992"
RESULT_URL = f"{BASE_URL}/result.aspx"

START_ROLL = 240505010001
END_ROLL = 240505010075
SKIP_ROLLS = set()

EXTRA_ROLLS = [
    240501010267, 
    240507010176
] + list(range(2520505001, 2520505007))

OUTPUT_FILE = "HPTU_Results_CSE_AIDS_4th_Sem.xlsx"
EXAM_TITLE = "B.TECH CSE (AI-DS) 4th SEMESTER REGULAR & RE-APPEAR (NEP) - SESSION MAY 2026"


def get_session_tokens(session):
    """Fetch the query page to get ASP.NET session tokens."""
    resp = session.get(QUERY_URL)
    resp.raise_for_status()
    soup = BeautifulSoup(resp.text, "html.parser")

    viewstate = soup.find("input", {"name": "__VIEWSTATE"})
    viewstate_gen = soup.find("input", {"name": "__VIEWSTATEGENERATOR"})
    hidden_id = soup.find("input", {"name": "id"})

    tokens = {}
    if viewstate:
        tokens["__VIEWSTATE"] = viewstate.get("value", "")
    if viewstate_gen:
        tokens["__VIEWSTATEGENERATOR"] = viewstate_gen.get("value", "")
    if hidden_id:
        tokens["id"] = hidden_id.get("value", "")

    return tokens


def fetch_result(session, roll_no, tokens):
    """Fetch the result for a given roll number."""
    form_data = {
        "__VIEWSTATE": tokens.get("__VIEWSTATE", ""),
        "__VIEWSTATEGENERATOR": tokens.get("__VIEWSTATEGENERATOR", ""),
        "RollNo": str(roll_no),
        "txtName": "",
        "id": tokens.get("id", "")
    }

    headers = {
        "Referer": QUERY_URL,
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
    }

    resp = session.post(RESULT_URL, data=form_data, headers=headers)
    resp.raise_for_status()
    return resp.text


def parse_result(html, roll_no):
    """Parse the result HTML and extract student data."""
    soup = BeautifulSoup(html, "html.parser")

    # Check if result exists
    page_text = soup.get_text()
    if "Roll No Not Exist" in page_text or "not found" in page_text.lower():
        return None

    structured = _parse_structured(html, roll_no)
    if structured:
        return structured

    # Fallback: generic heuristic parsing for unknown layouts
    result = {"Roll No": str(roll_no)}

    # Extract personal details from the result page
    # Look for text patterns in the page
    all_text = soup.get_text(separator="\n")
    lines = [line.strip() for line in all_text.split("\n") if line.strip()]

    # Try to find student name and father's name
    for i, line in enumerate(lines):
        if "Student Name" in line or "student name" in line.lower():
            # Name might be on the same line after colon or on the next line
            if ":" in line:
                result["Student Name"] = line.split(":", 1)[1].strip()
            elif i + 1 < len(lines):
                result["Student Name"] = lines[i + 1].strip()
        elif "Father" in line:
            if ":" in line:
                result["Father's Name"] = line.split(":", 1)[1].strip()
            elif i + 1 < len(lines):
                result["Father's Name"] = lines[i + 1].strip()

    # Try finding details in table rows
    tables = soup.find_all("table")
    for table in tables:
        rows = table.find_all("tr")
        for row in rows:
            cells = row.find_all(["td", "th"])
            cell_texts = [c.get_text(strip=True) for c in cells]

            # Look for personal detail rows
            for j, ct in enumerate(cell_texts):
                if "Roll No" in ct and j + 1 < len(cell_texts):
                    val = cell_texts[j + 1]
                    if val and val != "Roll No" and val not in [":", ""]:
                        result["Roll No"] = val
                elif "Student Name" in ct and j + 1 < len(cell_texts):
                    val = cell_texts[j + 1]
                    if val and val != "Student Name":
                        result["Student Name"] = val
                elif "Father" in ct and j + 1 < len(cell_texts):
                    val = cell_texts[j + 1]
                    if val and "Father" not in val:
                        result["Father's Name"] = val

    # Extract subjects/marks table
    subjects = []
    for table in tables:
        headers = []
        header_row = table.find("tr")
        if header_row:
            headers = [th.get_text(strip=True).lower() for th in header_row.find_all(["th", "td"])]

        # Check if this table has subject-related headers
        has_subject = any("subject" in h for h in headers)
        has_grade = any("grade" in h for h in headers)
        has_credit = any("credit" in h for h in headers)

        if has_subject or has_grade:
            rows = table.find_all("tr")[1:]  # Skip header row
            for row in rows:
                cells = row.find_all("td")
                if len(cells) >= 3:
                    cell_texts = [c.get_text(strip=True) for c in cells]
                    subject_data = {}

                    for idx, header in enumerate(headers):
                        if idx < len(cell_texts):
                            if "subject" in header and "code" not in header:
                                subject_data["Subject"] = cell_texts[idx]
                            elif "code" in header:
                                subject_data["Code"] = cell_texts[idx]
                            elif "credit" in header:
                                subject_data["Credit"] = cell_texts[idx]
                            elif "grade" in header:
                                subject_data["Grade"] = cell_texts[idx]

                    if subject_data.get("Subject") or subject_data.get("Code"):
                        subjects.append(subject_data)

    result["Subjects"] = subjects

    # Extract SGPA, CGPA, Result
    for i, line in enumerate(lines):
        lower = line.lower()
        if "sgpa" in lower:
            if ":" in line:
                result["SGPA"] = line.split(":", 1)[1].strip()
            else:
                # Try next line or extract number from same line
                import re
                nums = re.findall(r"\d+\.?\d*", line)
                if nums:
                    result["SGPA"] = nums[-1]
                elif i + 1 < len(lines):
                    result["SGPA"] = lines[i + 1].strip()
        elif "cgpa" in lower:
            if ":" in line:
                result["CGPA"] = line.split(":", 1)[1].strip()
            else:
                import re
                nums = re.findall(r"\d+\.?\d*", line)
                if nums:
                    result["CGPA"] = nums[-1]
                elif i + 1 < len(lines):
                    result["CGPA"] = lines[i + 1].strip()

    # Look for SGPA/CGPA in tables
    for table in tables:
        for row in table.find_all("tr"):
            cells = row.find_all(["td", "th"])
            cell_texts = [c.get_text(strip=True) for c in cells]
            for j, ct in enumerate(cell_texts):
                if "SGPA" in ct and j + 1 < len(cell_texts):
                    val = cell_texts[j + 1]
                    if val and val != "SGPA":
                        result["SGPA"] = val
                    elif ":" in ct:
                        result["SGPA"] = ct.split(":", 1)[1].strip()
                elif "CGPA" in ct and j + 1 < len(cell_texts):
                    val = cell_texts[j + 1]
                    if val and val != "CGPA":
                        result["CGPA"] = val
                    elif ":" in ct:
                        result["CGPA"] = ct.split(":", 1)[1].strip()

    # Find result status (PASS/FAIL/RE-APPEAR)
    for line in lines:
        lower = line.lower().strip()
        if lower in ["pass", "fail", "re-appear", "reappear", "absent", "comp"]:
            result["Result"] = line.strip().upper()
        elif "result" in lower and ":" in line:
            val = line.split(":", 1)[1].strip()
            if val.upper() in ["PASS", "FAIL", "RE-APPEAR", "REAPPEAR", "ABSENT", "COMP"]:
                result["Result"] = val.upper()

    # Also check in tables for result status
    for table in tables:
        for row in table.find_all("tr"):
            cells = row.find_all(["td", "th"])
            cell_texts = [c.get_text(strip=True) for c in cells]
            for j, ct in enumerate(cell_texts):
                if ct.strip().lower() == "result" and j + 1 < len(cell_texts):
                    result["Result"] = cell_texts[j + 1].strip()

    return result


def create_excel(all_results):
    """Create a styled Excel workbook from the results."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = re.sub(r'[\\/*?:\[\]]', '', EXAM_TITLE)[:31] or "Results"

    # Collect all unique subjects across all students
    all_subjects = []
    seen_subjects = set()
    for r in all_results:
        for s in r.get("Subjects", []):
            key = s.get("Code", s.get("Subject", ""))
            if key and key not in seen_subjects:
                seen_subjects.add(key)
                all_subjects.append(s)

    # Styles
    header_font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    pass_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    fail_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    grade_r_fill = PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid")  # Light red for Grade R

    # Title row
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6 + len(all_subjects))
    title_cell = ws.cell(row=1, column=1)
    title_cell.value = f"HIMACHAL PRADESH TECHNICAL UNIVERSITY\n{EXAM_TITLE.upper()}"
    title_cell.font = Font(name="Calibri", bold=True, size=14, color="1F3864")
    title_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 45

    # Build headers
    headers = ["S.No", "Roll No", "Student Name", "Father's Name"]
    for sub in all_subjects:
        code = sub.get("Code", sub.get("Subject", "Unknown"))
        name = sub.get("Subject", "")
        header_text = f"{code}" if code else name
        headers.append(header_text)
    headers.extend(["SGPA", "CGPA", "Result"])

    # Write header row
    row_num = 3
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    ws.row_dimensions[row_num].height = 35

    # Subject name sub-headers
    sub_header_row = row_num + 1
    for col in range(1, 5):
        cell = ws.cell(row=sub_header_row, column=col, value="")
        cell.border = thin_border
    for idx, sub in enumerate(all_subjects):
        col = 5 + idx
        name = sub.get("Subject", "")
        credit = sub.get("Credit", "")
        cell = ws.cell(row=sub_header_row, column=col, value=f"{name}\n(Credit: {credit})" if name else "")
        cell.font = Font(name="Calibri", size=8, italic=True, color="2F5496")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    for col in range(5 + len(all_subjects), 5 + len(all_subjects) + 3):
        cell = ws.cell(row=sub_header_row, column=col, value="")
        cell.border = thin_border

    ws.row_dimensions[sub_header_row].height = 50

    # Write data rows
    data_start = sub_header_row + 1
    for i, result in enumerate(all_results):
        row = data_start + i
        ws.cell(row=row, column=1, value=i + 1).alignment = cell_align
        ws.cell(row=row, column=2, value=result.get("Roll No", "")).alignment = cell_align
        ws.cell(row=row, column=3, value=result.get("Student Name", "")).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=row, column=4, value=result.get("Father's Name", "")).alignment = Alignment(horizontal="left", vertical="center")

        # Map subject grades
        student_subjects = {s.get("Code", s.get("Subject", "")): s for s in result.get("Subjects", [])}
        for idx, sub in enumerate(all_subjects):
            key = sub.get("Code", sub.get("Subject", ""))
            grade = student_subjects.get(key, {}).get("Grade", "-")
            cell = ws.cell(row=row, column=5 + idx, value=grade)
            cell.alignment = cell_align
            if grade.strip().upper() == "R":
                cell.fill = grade_r_fill

        # SGPA, CGPA, Result
        sgpa_col = 5 + len(all_subjects)
        cgpa_col = sgpa_col + 1
        result_col = cgpa_col + 1

        sgpa_val = result.get("SGPA", "")
        cgpa_val = result.get("CGPA", "")
        try:
            sgpa_val = float(sgpa_val)
        except (ValueError, TypeError):
            pass
        try:
            cgpa_val = float(cgpa_val)
        except (ValueError, TypeError):
            pass

        ws.cell(row=row, column=sgpa_col, value=sgpa_val).alignment = cell_align
        ws.cell(row=row, column=cgpa_col, value=cgpa_val).alignment = cell_align

        result_cell = ws.cell(row=row, column=result_col, value=result.get("Result", ""))
        result_cell.alignment = cell_align
        result_cell.font = Font(bold=True)

        if result.get("Result", "").upper() == "PASS":
            result_cell.fill = pass_fill
        elif result.get("Result", "").upper() in ["FAIL", "RE-APPEAR", "REAPPEAR"]:
            result_cell.fill = fail_fill

        # Apply borders to all cells in this row
        for col in range(1, result_col + 1):
            ws.cell(row=row, column=col).border = thin_border

    # Set column widths
    ws.column_dimensions["A"].width = 6   # S.No
    ws.column_dimensions["B"].width = 18  # Roll No
    ws.column_dimensions["C"].width = 25  # Student Name
    ws.column_dimensions["D"].width = 25  # Father's Name
    for idx in range(len(all_subjects)):
        col_letter = openpyxl.utils.get_column_letter(5 + idx)
        ws.column_dimensions[col_letter].width = 14
    sgpa_letter = openpyxl.utils.get_column_letter(5 + len(all_subjects))
    cgpa_letter = openpyxl.utils.get_column_letter(6 + len(all_subjects))
    result_letter = openpyxl.utils.get_column_letter(7 + len(all_subjects))
    ws.column_dimensions[sgpa_letter].width = 10
    ws.column_dimensions[cgpa_letter].width = 10
    ws.column_dimensions[result_letter].width = 14

    # Freeze panes
    ws.freeze_panes = f"E{data_start}"

    wb.save(OUTPUT_FILE)
    print(f"\n{'='*60}")
    print(f"  Excel file saved: {OUTPUT_FILE}")
    print(f"  Total students: {len(all_results)}")
    print(f"{'='*60}")


def parse_args():
    global QUERY_URL, RESULT_URL, OUTPUT_FILE
    p = argparse.ArgumentParser(description="HPTU bulk result scraper -> Excel")
    p.add_argument("--url", default=QUERY_URL, help="query.aspx result URL")
    p.add_argument("--start", type=int, default=START_ROLL, help="first roll number")
    p.add_argument("--end", type=int, default=END_ROLL, help="last roll number")
    p.add_argument("--skip", type=int, nargs="*", default=sorted(SKIP_ROLLS), help="roll numbers to skip")
    p.add_argument("--extra", type=int, nargs="*", default=EXTRA_ROLLS, help="additional roll numbers")
    p.add_argument("--out", default=OUTPUT_FILE, help="output .xlsx file")
    args = p.parse_args()
    QUERY_URL = args.url
    RESULT_URL = args.url.split("?")[0].replace("query.aspx", "result.aspx")
    OUTPUT_FILE = args.out
    return args


def detect_exam_title(session):
    """Read the exam name from the query page <title> (e.g. 'HIMTU: B. Tech ...')."""
    global EXAM_TITLE
    try:
        soup = BeautifulSoup(session.get(QUERY_URL, timeout=15).text, "html.parser")
        title = soup.title.get_text(strip=True) if soup.title else ""
        title = re.sub(r"^\s*HIMTU\s*:\s*", "", title)
        if title:
            EXAM_TITLE = title
    except Exception:
        pass


def main():
    args = parse_args()
    session = requests.Session()
    detect_exam_title(session)

    print("=" * 60)
    print("  HPTU Result Scraper")
    print(f"  {EXAM_TITLE}")
    print("=" * 60)

    # Step 1: Get session tokens
    print("\n[*] Fetching session tokens...")
    try:
        tokens = get_session_tokens(session)
        print(f"[+] Session tokens acquired ({len(tokens)} tokens)")
    except Exception as e:
        print(f"[!] Failed to get session tokens: {e}")
        sys.exit(1)

    all_results = []
    failed_rolls = []
    not_found_rolls = []

    # Build the full list of roll numbers to process
    skip = set(args.skip or [])
    rolls_to_process = [r for r in range(args.start, args.end + 1) if r not in skip]
    rolls_to_process.extend(r for r in (args.extra or []) if r not in skip and r not in rolls_to_process)
    
    total = len(rolls_to_process)
    current = 0

    for roll_no in rolls_to_process:
        current += 1
        roll_str = str(roll_no)
        print(f"\n[{current}/{total}] Fetching result for Roll No: {roll_str}...", end=" ")

        try:
            html = fetch_result(session, roll_str, tokens)
            result = parse_result(html, roll_str)

            if result:
                all_results.append(result)
                name = result.get("Student Name", "N/A")
                sgpa = result.get("SGPA", "N/A")
                status = result.get("Result", "N/A")
                print(f"✓ {name} | SGPA: {sgpa} | {status}")
            else:
                not_found_rolls.append(roll_str)
                print("✗ Roll No Not Found")

            # Refresh tokens periodically to avoid session expiry
            if current % 15 == 0:
                try:
                    tokens = get_session_tokens(session)
                    print("[*] Session tokens refreshed")
                except Exception:
                    pass

        except Exception as e:
            failed_rolls.append(roll_str)
            print(f"✗ Error: {e}")

        # Small delay to be respectful to the server
        time.sleep(0.5)

    # Summary
    print(f"\n{'='*60}")
    print(f"  Scraping Complete!")
    print(f"  Successful: {len(all_results)}")
    print(f"  Not Found:  {len(not_found_rolls)}")
    print(f"  Failed:     {len(failed_rolls)}")
    print(f"{'='*60}")

    if not_found_rolls:
        print(f"\n[!] Roll numbers not found: {', '.join(not_found_rolls)}")
    if failed_rolls:
        print(f"[!] Failed roll numbers: {', '.join(failed_rolls)}")

    # Create Excel
    if all_results:
        print("\n[*] Creating Excel spreadsheet...")
        create_excel(all_results)
    else:
        print("\n[!] No results to save.")


if __name__ == "__main__":
    main()
